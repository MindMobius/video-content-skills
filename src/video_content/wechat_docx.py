"""Canonical, locally validated DOCX and transient clipboard transports.

Content is the manuscript; a successful local check is not a WeChat import receipt.
"""

from __future__ import annotations

import base64
import hashlib
import io
import os
import posixpath
import re
import tempfile
import zipfile
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from .frames import image_dimensions
from .store import Store
from .wechat_renderer import DEFAULT_DISCLAIMER

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PIC = "http://schemas.openxmlformats.org/drawingml/2006/picture"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
PACKAGE_PROFILE = "python-docx-1.2.0"


def text_digest(text: str) -> str:
    # Ignore layout whitespace only; punctuation and every substantive character count.
    return hashlib.sha256(re.sub(r"[\s\u200b\ufeff]+", "", text).encode()).hexdigest()


def manuscript_parts(document: dict[str, Any]) -> list[dict[str, Any]]:
    blocks = document["blocks"]
    source = document["source"]
    parts: list[dict[str, Any]] = []
    if not any(
        str(b.get("text", "")).startswith("原始内容：")
        and source["canonical_url"] in str(b.get("text", ""))
        for b in blocks
    ):
        parts.append(
            {
                "text": f"原始内容：{source['title']}｜创作者：{source['creator']}｜{source['canonical_url']}"
            }
        )
        parts.append({"text": source.get("disclaimer") or DEFAULT_DISCLAIMER})
    for block in blocks:
        kind = block["type"]
        if kind == "image":
            parts.append(block)
            if block.get("caption"):
                parts.append({"text": block["caption"]})
        elif kind == "list":
            parts.extend({"text": item} for item in block["items"])
        elif kind != "separator":
            parts.append({"text": block["text"], "heading": kind == "heading"})
    return parts


def _image(store: Store, job_id: str, artifact_id: str) -> tuple[dict[str, Any], bytes]:
    ref, data = store.read_artifact(job_id, artifact_id)
    digest = hashlib.sha256(data).hexdigest()
    if digest != ref["sha256"]:
        raise ValueError("Canonical image bytes changed")
    if not data.startswith((b"\x89PNG", b"\xff\xd8")):
        raise ValueError(
            "DOCX transport supports verified PNG/JPEG bytes; do not omit the image"
        )
    width, height = image_dimensions(store.job_dir(job_id) / ref["path"])
    return {
        "artifact_id": artifact_id,
        "sha256": digest,
        "pixel_width": width,
        "pixel_height": height,
    }, data


def _audit_package(
    store: Store, job_id: str, content: dict[str, Any], payload: bytes
) -> dict[str, Any]:
    """Independently read package bytes, including each image's position in text.

    This intentionally accepts the previous minimal package for resume only;
    its original checkpoint hash must be checked by the caller first.
    """
    expected_images, expected_segments, text_buffer = [], [], []
    for part in manuscript_parts(content["document"]):
        if "artifact_id" in part:
            expected_segments.append(text_digest("".join(text_buffer)))
            text_buffer = []
            expected_images.append(_image(store, job_id, part["artifact_id"])[0])
        else:
            text_buffer.append(str(part["text"]))
    expected_segments.append(text_digest("".join(text_buffer)))
    expected_text = "".join(
        str(p["text"])
        for p in manuscript_parts(content["document"])
        if "artifact_id" not in p
    )
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or archive.testzip() is not None:
            raise ValueError("Invalid DOCX ZIP package")
        for name in names:
            if name.endswith(".rels"):
                for rel in ET.fromstring(archive.read(name)):
                    if rel.get("TargetMode") == "External":
                        raise ValueError("DOCX must not contain external relationships")
        root = ET.fromstring(archive.read("word/document.xml"))
        relations = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
        rels = {r.attrib["Id"]: r for r in relations}
        if len(rels) != len(relations):
            raise ValueError("Duplicate DOCX relationship identity")
        texts, segments, text_buffer, embedded = [], [], [], []
        for node in root.iter():
            if node.tag == f"{{{W}}}t":
                texts.append(node.text or "")
                text_buffer.append(node.text or "")
            elif node.tag == f"{{{A}}}blip":
                segments.append(text_digest("".join(text_buffer)))
                text_buffer = []
                embedded.append(node)
        segments.append(text_digest("".join(text_buffer)))
        if text_digest("".join(texts)) != text_digest(expected_text):
            raise ValueError("DOCX text drift")
        if len(embedded) != len(expected_images):
            raise ValueError("DOCX embedded image order/count drift")
        if segments != expected_segments:
            raise ValueError("DOCX image position drift relative to manuscript text")
        for blip, image in zip(embedded, expected_images):
            rel = rels.get(blip.get(f"{{{R}}}embed"))
            if rel is None or rel.get("Type") != f"{R}/image":
                raise ValueError("Invalid DOCX image relationship")
            target = posixpath.normpath("word/" + rel.attrib["Target"])
            if not target.startswith("word/media/") or "\\" in target:
                raise ValueError("DOCX image path escapes embedded media")
            if hashlib.sha256(archive.read(target)).hexdigest() != image["sha256"]:
                raise ValueError("DOCX embedded image bytes/order drift")
        drawings = list(root.iter(f"{{{WP}}}inline"))
        if len(drawings) != len(expected_images) or list(root.iter(f"{{{WP}}}anchor")):
            raise ValueError("DOCX image count/inline placement drift")
        for drawing, image in zip(drawings, expected_images):
            extent = drawing.find(f"{{{WP}}}extent")
            shape = drawing.find(f".//{{{A}}}xfrm/{{{A}}}ext")
            for geometry in (extent, shape):
                if geometry is None:
                    raise ValueError("Missing DOCX image geometry")
                cx, cy = int(geometry.attrib["cx"]), int(geometry.attrib["cy"])
                if (
                    min(cx, cy) <= 0
                    or abs(
                        (cx / cy) / (image["pixel_width"] / image["pixel_height"]) - 1
                    )
                    > 0.001
                ):
                    raise ValueError("DOCX image aspect drift")
            if extent.attrib != shape.attrib:
                raise ValueError("DOCX image geometry disagrees across drawing extents")
    return {
        "sha256": hashlib.sha256(payload).hexdigest(),
        "text_sha256": text_digest(expected_text),
        "images": expected_images,
        "validated": True,
        "validation_scope": "local_package_only",
    }


def inspect_bound_docx(
    store: Store,
    job_id: str,
    content: dict[str, Any],
    path: Path,
    *,
    expected_sha256: str,
) -> dict[str, Any]:
    """Resume exact already-bound bytes without invoking a newer generator."""
    if not path.is_file():
        raise ValueError(
            "Bound DOCX is missing; preserve checkpoint and recover, do not regenerate"
        )
    payload = path.read_bytes()
    if hashlib.sha256(payload).hexdigest() != expected_sha256:
        raise ValueError("Bound DOCX bytes changed; preserve checkpoint and diagnose")
    return {"path": str(path), **_audit_package(store, job_id, content, payload)}


def build_docx(
    store: Store, job_id: str, content: dict[str, Any], path: Path
) -> dict[str, Any]:
    # The distribution pins this generator. No handwritten Word template/theme.
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Emu, Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError(
            "DOCX generation requires project dependencies; run python scripts/bootstrap.py --apply"
        ) from exc

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)
    # Retain the same pixel aspect even when files have asymmetric DPI metadata.
    max_width = int(Inches(6))
    max_height = int(Inches(8.5))
    for style_id, size in (("Normal", 11), ("Heading 2", 13)):
        style = document.styles[style_id]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei"
        )
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.5
    for part in manuscript_parts(content["document"]):
        if "artifact_id" in part:
            image, data = _image(store, job_id, part["artifact_id"])
            scale = min(
                max_width / image["pixel_width"], max_height / image["pixel_height"]
            )
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.add_run().add_picture(
                io.BytesIO(data),
                width=Emu(round(scale * image["pixel_width"])),
                height=Emu(round(scale * image["pixel_height"])),
            )
        else:
            document.add_paragraph(
                str(part["text"]),
                style="Heading 2" if part.get("heading") else "Normal",
            )
    properties = document.core_properties
    properties.title = content["document"]["title"]
    properties.author = "video-content-skills"
    properties.last_modified_by = "video-content-skills"
    properties.identifier = content["content_id"]
    properties.comments = "Canonical Content transport; not proof of WeChat import."
    created = content.get("created_at")
    if created:
        instant = datetime.fromisoformat(created.replace("Z", "+00:00"))
        properties.created = properties.modified = instant
    buffer = io.BytesIO()
    document.save(buffer)
    # ZIP timestamps are not manuscript data. Fix them for repeatable artifact hashes.
    output = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(output, "w") as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "docProps/app.xml":
                app = ET.fromstring(data)
                ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
                app.find(
                    f"{{{ns}}}Application"
                ).text = f"video-content-skills / {PACKAGE_PROFILE}"
                company = app.find(f"{{{ns}}}Company")
                if company is not None:
                    company.text = ""
                data = ET.tostring(app, encoding="utf-8", xml_declaration=True)
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)
    payload = output.getvalue()
    report = _audit_package(store, job_id, content, payload)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if path.read_bytes() != payload:
            raise ValueError(
                "Existing DOCX differs from canonical Content; preserve it for diagnosis before regenerating"
            )
    else:
        descriptor, temporary = tempfile.mkstemp(
            prefix=".docx-", suffix=".tmp", dir=path.parent
        )
        try:
            with os.fdopen(descriptor, "wb") as stream:
                stream.write(payload)
                stream.flush()
                os.fsync(stream.fileno())
            # Never overwrite a concurrently created transport.
            os.link(temporary, path)
        finally:
            Path(temporary).unlink(missing_ok=True)
    return {"path": str(path), **report}


def build_canonical_clipboard(
    store: Store, job_id: str, content: dict[str, Any]
) -> dict[str, Any]:
    """An in-memory alternative transport of the exact same manuscript parts."""
    html_parts, texts, image_ids = [], [], []
    for part in manuscript_parts(content["document"]):
        if "artifact_id" in part:
            ref, data = store.read_artifact(job_id, part["artifact_id"])
            mime = (
                "image/png"
                if data.startswith(b"\x89PNG")
                else "image/jpeg"
                if data.startswith(b"\xff\xd8")
                else None
            )
            if mime is None:
                raise ValueError("Unsupported canonical image bytes")
            if hashlib.sha256(data).hexdigest() != ref["sha256"]:
                raise ValueError("Canonical image bytes changed")
            image_ids.append(part["artifact_id"])
            encoded = base64.b64encode(data).decode("ascii")
            html_parts.append(
                f'<p><img src="data:{mime};base64,{encoded}" style="max-width:100%;height:auto" /></p>'
            )
        else:
            text = str(part["text"])
            texts.append(text)
            tag = "h2" if part.get("heading") else "p"
            html_parts.append(f"<{tag}>{escape(text)}</{tag}>")
    html = "".join(html_parts)
    return {
        "html": html,
        "plain_text": "\n".join(texts),
        "text_sha256": text_digest("".join(texts)),
        "image_ids": image_ids,
        "sha256": hashlib.sha256(html.encode()).hexdigest(),
    }
